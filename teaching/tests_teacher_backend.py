"""教师导入、看板、成绩、提交和批阅的真实 API 权限与数据回归。"""
import csv
import io
import json
import zipfile
from datetime import timedelta

from django.contrib.auth.models import Group
from django.db import IntegrityError, transaction
from django.test.utils import CaptureQueriesContext
from django.db import connection
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from teaching.tests import ObjectPermissionTestBase
from teaching.tests_problem_import import buildMd
from teaching.models import Homework, HomeworkProblem, Problem, Submission, SubmissionReview
from users.models import User


class TeacherBackendTests(ObjectPermissionTestBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.problem = Problem.objects.create(
            title='题一', owner=cls.teacher, languages=['python', 'c'],
            tests=[{'input': 'HIDDEN_INPUT', 'output': 'HIDDEN_EXPECTED', 'points': 100}])
        cls.problem2 = Problem.objects.create(
            title='题二', owner=cls.teacher, languages=['python'],
            tests=[{'input': 'HIDDEN_INPUT_2', 'output': 'HIDDEN_EXPECTED_2', 'points': 100}])
        cls.homework = Homework.objects.create(
            teaching_class=cls.myTeachingClass, title='两题作业', created_by=cls.teacher,
            start_at=timezone.now() - timedelta(hours=1),
            deadline=timezone.now() + timedelta(hours=1))
        for order, problem in enumerate((cls.problem, cls.problem2), 1):
            HomeworkProblem.objects.create(homework=cls.homework, problem=problem,
                                            order=order, problem_revision=1)
        cls.unsubmitted = User.objects.create_user('S9013', name='未交学生')
        cls.unsubmitted.groups.add(Group.objects.get(name='学生'))
        cls._makeStudent(cls.unsubmitted, '20239003', cls.myProfile.major, cls.myAdminClass)

    def setUp(self):
        self.client = APIClient()
        self.client.force_login(self.teacher)
        self.classUrl = f'/api/teacher/teaching-classes/{self.myTeachingClass.pk}'
        self.gradeUrl = f'/api/teacher/homeworks/{self.homework.pk}/grades/'

    def submit(self, score=100, status=Submission.STATUS_ACCEPTED, problem=None, user=None):
        problem = problem or self.problem
        return Submission.objects.create(
            user=user or self.student, homework=self.homework, problem=problem,
            submit_type=Submission.SUBMIT_HOMEWORK, language='python', code='print(42)',
            fingerprint='test', status=status, score=score, total_score=100,
            problem_revision=1, tests_hash='test', judge_snapshot={'tests': problem.tests})

    def zipFile(self, code='A001', title='导入题', **files):
        from django.core.files.uploadedfile import SimpleUploadedFile
        buffer = io.BytesIO()
        rows = io.StringIO()
        csv.writer(rows).writerows([['file', 'code'], ['a.md', code]])
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('manifest.csv', rows.getvalue())
            archive.writestr('a.md', buildMd(title=title))
            for name, content in files.items():
                archive.writestr(name, content)
        return SimpleUploadedFile('questions.zip', buffer.getvalue(), content_type='application/zip')

    def importZip(self, **kwargs):
        return self.client.post('/api/teacher/problems/import/',
                                {'file': self.zipFile(**kwargs)}, format='multipart')

    def test_personal_import_and_history_are_owned(self):
        response = self.importZip()
        self.assertEqual(response.status_code, 200, response.data)
        imported = Problem.objects.get(source_code='A001')
        self.assertFalse(imported.is_public)
        self.assertEqual(imported.owner, self.teacher)
        self.assertEqual(self.importZip().data['skippedCount'], 1)
        history = self.client.get('/api/teacher/problems/imports/').json()
        self.assertEqual(history['count'], 2)
        self.client.force_login(self.teacher2)
        self.assertEqual(self.client.get('/api/teacher/problems/imports/').json()['count'], 0)

    def test_same_code_in_separate_libraries_never_overwrites(self):
        self.importZip()
        original = Problem.objects.get(source_code='A001')
        self.client.force_login(self.teacher2)
        self.assertEqual(self.importZip(title='另一位教师').data['createdCount'], 1)
        self.client.force_login(self.superUser)
        response = self.client.post('/api/admin/problems/import/',
                                    {'file': self.zipFile(title='公共题')}, format='multipart')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['createdCount'], 1)
        original.refresh_from_db()
        self.assertEqual(original.title, '导入题')
        self.assertEqual(Problem.objects.filter(source_code='A001').count(), 3)

    def test_private_overwrite_changes_revision_and_keeps_homework_snapshot(self):
        self.importZip()
        problem = Problem.objects.get(source_code='A001')
        item = HomeworkProblem.objects.create(
            homework=self.homework, problem=problem, order=3, problem_revision=1)
        response = self.client.post('/api/teacher/problems/import/', {
            'file': self.zipFile(title='修改后'), 'overwrite': 'true'}, format='multipart')
        self.assertEqual(response.data['updatedCount'], 1)
        problem.refresh_from_db()
        item.refresh_from_db()
        self.assertEqual(problem.revision, 2)
        self.assertEqual(item.problem_snapshot['title'], '导入题')

    def test_failed_import_is_visible_in_own_history(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        response = self.client.post('/api/teacher/problems/import/',
            {'file': SimpleUploadedFile('bad.zip', b'bad')}, format='multipart')
        self.assertEqual(response.status_code, 400)
        history = self.client.get('/api/teacher/problems/imports/').json()
        self.assertEqual(history['count'], 1)
        self.assertEqual(history['results'][0]['status'], 'failed')
        self.assertTrue(history['results'][0]['error'])

    def test_import_rejects_forged_owner_and_non_teachers(self):
        response = self.client.post('/api/teacher/problems/import/', {
            'file': self.zipFile(), 'ownerId': self.teacher2.pk, 'isPublic': 'true'}, format='multipart')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Problem.objects.filter(source_code='A001').exists())
        for user in (self.student, self.courseTeam):
            self.client.force_login(user)
            self.assertEqual(self.importZip().status_code, 403)

    def test_database_enforces_code_within_scope(self):
        for public, owner in ((True, self.superUser), (False, self.teacher)):
            Problem.objects.create(title='一', source_code='SAME', is_public=public, owner=owner)
            with self.assertRaises(IntegrityError), transaction.atomic():
                Problem.objects.create(title='二', source_code='SAME', is_public=public, owner=owner)

    def test_grades_sum_per_problem_best_and_exclude_system_pending_repeat(self):
        self.submit(70, Submission.STATUS_WRONG_ANSWER)
        self.submit(100)
        self.submit(80, Submission.STATUS_WRONG_ANSWER, problem=self.problem2)
        self.submit(99, Submission.STATUS_SYSTEM_ERROR, problem=self.problem2)
        self.submit(99, Submission.STATUS_REPEAT, problem=self.problem2)
        self.submit(100, Submission.STATUS_RUNNING, problem=self.problem2)
        response = self.client.get(self.gradeUrl).json()
        student = next(r for r in response['results'] if r['id'] == self.student.pk)
        self.assertEqual(student['score'], 180)
        self.assertEqual(student['totalScore'], 200)
        self.assertEqual(student['pendingCount'], 1)
        self.assertTrue(student['gradingComplete'])
        self.assertEqual(response['summary']['averageScore'], 180)
        self.client.force_login(self.student)
        row = next(r for r in self.client.get('/api/homeworks/').json()
                   if r['id'] == self.homework.pk)
        self.assertEqual(row['myBestScore'], 180)

    def test_unsubmitted_and_system_error_remain_distinct(self):
        self.submit(100, Submission.STATUS_SYSTEM_ERROR)
        data = self.client.get(self.gradeUrl).json()
        submitted = next(r for r in data['results'] if r['id'] == self.student.pk)
        absent = next(r for r in data['results'] if r['id'] == self.unsubmitted.pk)
        self.assertIsNone(submitted['score'])
        self.assertEqual(submitted['problems'][0]['state'], 'system_error')
        self.assertEqual(absent['submissionState'], 'not_submitted')
        self.assertIsNone(absent['score'])
        self.assertIsNone(data['summary']['averageScore'])
        self.assertEqual(self.client.get(self.gradeUrl, {
            'submissionState': 'not_submitted'}).json()['count'], 1)

    def test_dashboard_students_and_detail_use_current_roster(self):
        self.submit()
        self.submit(problem=self.problem2)
        data = self.client.get(self.classUrl + '/dashboard/').json()
        self.assertEqual(data['studentCount'], 2)
        self.assertEqual(data['homeworks'][0]['passedStudentCount'], 1)
        self.assertNotIn('students', data)
        students = self.client.get(self.classUrl + '/students/', {'q': 'S9011'}).json()
        self.assertEqual(students['count'], 1)
        detail = self.client.get(self.classUrl + f'/students/{self.student.pk}/').json()
        self.assertEqual(detail['homeworks'][0]['score'], 200)
        self.assertEqual(self.client.get(self.classUrl + f'/students/{self.student2.pk}/').status_code, 404)

    def test_board_queries_do_not_scale_per_student(self):
        with CaptureQueriesContext(connection) as before:
            self.client.get(self.classUrl + '/dashboard/')
        for i in range(8):
            user = User.objects.create_user(f'bulk-{i}')
            self._makeStudent(user, f'bulk-{i}', self.myProfile.major, self.myAdminClass)
        with CaptureQueriesContext(connection) as after:
            self.client.get(self.classUrl + '/dashboard/')
        self.assertLessEqual(len(after), len(before) + 1)

    def test_class_submission_list_filters_and_detail_whitelist(self):
        submission = self.submit()
        response = self.client.get(self.classUrl + '/submissions/', {
            'student': self.student.pk, 'homework': self.homework.pk, 'status': 'accepted'}).json()
        self.assertEqual(response['count'], 1)
        self.assertNotIn('code', response['results'][0])
        detail = self.client.get(f'/api/teacher/submissions/{submission.pk}/').json()
        self.assertEqual(detail['code'], 'print(42)')
        for key in ('tests', 'judge_snapshot', 'spj_code', 'file_id'):
            self.assertNotIn(key, detail)
        self.assertNotIn('HIDDEN_', str(detail))

    def test_teachers_cannot_cross_class_boundary_and_students_cannot_use_teacher_api(self):
        submission = self.submit()
        for user, expected in ((self.teacher2, 404), (self.student, 403),
                               (self.courseTeam, 403)):
            self.client.force_login(user)
            for url in (self.classUrl + '/dashboard/', self.classUrl + '/students/',
                        self.classUrl + '/submissions/', self.gradeUrl,
                        self.gradeUrl + 'export/', f'/api/teacher/submissions/{submission.pk}/'):
                self.assertEqual(self.client.get(url).status_code, expected, url)

    def test_departed_member_history_survives_without_unsubmitted_roster_entry(self):
        submission = self.submit()
        member = self.myProfile.memberships.get(left_at__isnull=True)
        member.left_at = timezone.now().date()
        member.save()
        self.assertEqual(self.client.get(self.gradeUrl).json()['count'], 1)
        self.assertEqual(self.client.get(self.classUrl + '/submissions/').json()['count'], 1)
        self.assertEqual(self.client.get(f'/api/teacher/submissions/{submission.pk}/').status_code, 200)
        self.assertEqual(self.client.get(self.classUrl + f'/students/{self.student.pk}/').status_code, 404)

    def test_export_matches_grades_and_text_never_becomes_formula(self):
        self.student.name = '=1+1'
        self.student.save(update_fields=['name'])
        self.submit()
        self.submit(60, Submission.STATUS_WRONG_ANSWER, problem=self.problem2)
        response = self.client.get(self.gradeUrl + 'export/')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.content)).active
        row = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == self.student.username)
        self.assertEqual(row[1].value, '=1+1')
        self.assertEqual(row[1].data_type, 's')
        self.assertEqual([c.value for c in row[4:8]], [100, 60, 160, 200])

    def test_review_create_update_delete_and_student_own_feedback(self):
        submission = self.submit()
        url = f'/api/teacher/submissions/{submission.pk}/review/'
        self.assertEqual(self.client.get(url).json(), {'review': None})
        response = self.client.put(url, {'comment': '请补充边界条件。'}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['review']['reviewedById'], self.teacher.pk)
        self.assertEqual(self.client.get(self.classUrl + '/submissions/', {
            'reviewed': 'true'}).json()['count'], 1)
        self.client.force_login(self.student)
        detail = self.client.get(f'/api/submissions/{submission.pk}/').json()
        self.assertEqual(detail['teacherReview']['comment'], '请补充边界条件。')
        self.client.force_login(self.student2)
        self.assertEqual(self.client.get(f'/api/submissions/{submission.pk}/').status_code, 404)
        self.client.force_login(self.teacher)
        self.client.put(url, {'comment': '已改进。'}, format='json')
        self.assertEqual(SubmissionReview.objects.count(), 1)
        submission.refresh_from_db()
        self.assertEqual((submission.status, submission.score), ('accepted', 100))
        self.assertEqual(self.client.delete(url).status_code, 204)
        self.assertFalse(SubmissionReview.objects.exists())

    def test_review_write_checks_permission_scope_and_payload(self):
        submission = self.submit()
        url = f'/api/teacher/submissions/{submission.pk}/review/'
        for data in ({'comment': ''}, {'comment': 1}, {'comment': 'x' * 5001},
                     {'comment': '有效', 'score': 0}, [], {'reviewedById': self.teacher.pk}):
            self.assertEqual(self.client.put(url, data, format='json').status_code, 400)
        self.client.force_login(self.teacher2)
        self.assertEqual(self.client.put(url, {'comment': '越权'}, format='json').status_code, 404)
        self.client.force_login(self.student)
        self.assertEqual(self.client.put(url, {'comment': '越权'}, format='json').status_code, 403)
        self.teacher.user_permissions.clear()
        self.teacher.groups.remove(Group.objects.get(name='教师'))
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.put(url, {'comment': '撤权'}, format='json').status_code, 403)

    def test_filters_validate_instead_of_returning_500(self):
        for params in ({'student': 'bad'}, {'homework': '-1'}, {'status': 'bogus'},
                       {'reviewed': 'maybe'}, {'problem': str(2 ** 63)}):
            self.assertEqual(self.client.get(self.classUrl + '/submissions/', params).status_code, 400)
        self.assertEqual(self.client.get(self.gradeUrl, {'submissionState': 'bad'}).status_code, 400)

    def test_cross_class_and_missing_submission_have_identical_404(self):
        submission = self.submit()
        self.client.force_login(self.teacher2)
        existing = self.client.get(f'/api/teacher/submissions/{submission.pk}/')
        missing = self.client.get('/api/teacher/submissions/999999/')
        self.assertEqual(existing.status_code, 404)
        self.assertEqual(existing.json(), missing.json())

    def test_empty_filtered_export_keeps_all_problem_columns(self):
        response = self.client.get(self.gradeUrl + 'export/', {'q': '不存在的学生'})
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual([c.value for c in sheet[1]][4:6], ['题一', '题二'])
        self.assertEqual(sheet.max_row, 1)

    def test_invalid_import_body_and_review_unicode_return_400(self):
        for body in ([{}], [], 'bad'):
            self.assertEqual(self.client.post('/api/teacher/problems/import/',
                                             body, format='json').status_code, 400)
        submission = self.submit()
        response = self.client.put(f'/api/teacher/submissions/{submission.pk}/review/',
            json.dumps({'comment': '\ud800'}), content_type='application/json')
        self.assertEqual(response.status_code, 400)

    def test_coteacher_can_review_without_changing_automatic_score(self):
        submission = self.submit()
        self.myTeachingClass.teachers.add(self.teacher2)
        self.client.force_login(self.teacher2)
        response = self.client.put(f'/api/teacher/submissions/{submission.pk}/review/',
                                   {'comment': '同班教师评阅'}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['review']['reviewedById'], self.teacher2.pk)
        submission.refresh_from_db()
        self.assertEqual(submission.score, 100)

    def test_frozen_titles_and_total_survive_live_problem_changes(self):
        self.submit()
        self.problem.title = '后续新题名'
        self.problem.tests = [{'input': 'new', 'output': 'new', 'points': 1}]
        self.problem.save()
        data = self.client.get(self.gradeUrl).json()
        row = next(r for r in data['results'] if r['id'] == self.student.pk)
        self.assertEqual(row['problems'][0]['title'], '题一')
        self.assertEqual(row['totalScore'], 200)
