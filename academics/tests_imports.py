"""教务导入的文件边界、物理行号、行级回滚及任课权限。"""
from io import BytesIO
from unittest import mock
from zipfile import ZipFile

from django.contrib.messages import get_messages
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from django.test import SimpleTestCase
from openpyxl import Workbook

from academics import views
from academics.api import AdminClassDetailView, CourseCollegeView, ImportView
from academics.models import AdminClass, ClassMembership, Course, Department, Major
from academics.tests import OfficeApiTestBase
from users.models import SchoolIdentity, StudentProfile, TeacherProfile, User


def xlsx(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile('导入.xlsx', output.getvalue())


class XlsxParserTests(SimpleTestCase):
    def test_blank_rows_keep_physical_positions_and_only_first_row_is_header(self):
        rows = views.parseXlsxRows(xlsx([
            ['学号', '姓名', '初始密码'], [], ['S1', '甲', 'pw'], [], ['S2', '', 'pw']]))
        self.assertFalse(any(rows[0]))
        self.assertEqual(rows[1][:3], ['S1', '甲', 'pw'])
        self.assertFalse(any(rows[2]))
        self.assertEqual(rows[3][:3], ['S2', '', 'pw'])
        rows = views.parseXlsxRows(xlsx([[], ['S1', '甲', 'pw']]))
        self.assertEqual(rows, [['S1', '甲', 'pw']])

    def test_numeric_identifiers_and_password_cells_keep_existing_conversion(self):
        self.assertEqual(views.parseXlsxRows(xlsx([
            ['编号', '姓名', '密码'], [20231001.0, ' 学生 ', 123456]])),
            [['20231001', '学生', '123456']])

    def test_malformed_or_non_xlsx_file_has_safe_validation_error(self):
        emptyZip = BytesIO()
        with ZipFile(emptyZip, 'w') as archive:
            archive.writestr('private-path.txt', 'SECRET_DETAIL')
        for content in (b'', b'not a workbook', emptyZip.getvalue()):
            with self.subTest(content=content[:20]):
                with self.assertRaises(ValidationError) as error:
                    views.parseXlsxRows(SimpleUploadedFile('bad.xlsx', content))
                self.assertNotIn('private-path', str(error.exception))
                self.assertNotIn('SECRET_DETAIL', str(error.exception))

    def test_file_read_is_bounded_and_oversize_is_rejected(self):
        class BoundedStream(BytesIO):
            def read(stream, size=-1):
                self.assertEqual(size, 9)
                return super(BoundedStream, stream).read(size)
        with mock.patch.object(views, 'MAX_XLSX_BYTES', 8):
            with self.assertRaises(ValidationError):
                views.parseXlsxRows(BoundedStream(b'x' * 100))
            with mock.patch.object(views, 'load_workbook') as loader:
                with self.assertRaises(ValidationError):
                    views.parseXlsxRows(SimpleUploadedFile('large.xlsx', b'x' * 9))
                loader.assert_not_called()

    def test_uncompressed_limit_is_checked_before_openpyxl(self):
        uploaded = xlsx([['编号', '姓名'], ['x', 'y']])
        with mock.patch.object(views, 'MAX_XLSX_UNCOMPRESSED_BYTES', 1):
            with mock.patch.object(views, 'load_workbook') as loader:
                with self.assertRaises(ValidationError):
                    views.parseXlsxRows(uploaded)
                loader.assert_not_called()

    def test_extra_content_columns_rejected_but_empty_columns_allowed(self):
        with self.assertRaises(ValidationError):
            views.parseXlsxRows(xlsx([['编号', '姓名', '密码'], ['a', 'b', 'c', '', 'd']]))
        self.assertEqual(views.parseXlsxRows(xlsx([
            ['编号', '姓名', '密码'], ['a', 'b', 'c', '', '']])), [['a', 'b', 'c']])

    def test_physical_row_limit_includes_blank_rows(self):
        with mock.patch.object(views, 'MAX_IMPORT_ROWS', 2):
            self.assertEqual(len(views.parseXlsxRows(xlsx([
                ['编号', '姓名'], [], ['a', 'b']]))), 2)
            with self.assertRaises(ValidationError):
                views.parseXlsxRows(xlsx([['编号', '姓名'], [], [], ['a', 'b']]))

    def test_workbook_is_closed_on_success_and_validation_error(self):
        for rows in ([('编号', '姓名'), ('a', 'b')],
                     [('编号', '姓名'), ('a', 'b', 'c', 'extra')]):
            workbook = mock.Mock()
            workbook.active.iter_rows.return_value = iter(rows)
            with mock.patch.object(views, 'load_workbook', return_value=workbook):
                try:
                    views.parseXlsxRows(xlsx([['编号', '姓名']]))
                except ValidationError:
                    pass
            workbook.close.assert_called_once_with()

    def test_actual_rows_are_used_even_when_sheet_dimension_claims_only_header(self):
        original = xlsx([['编号', '姓名'], ['a', 'b']]).read()
        rewritten = BytesIO()
        with ZipFile(BytesIO(original)) as source, ZipFile(rewritten, 'w') as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    content = content.replace(b'ref="A1:B2"', b'ref="A1:A1"')
                target.writestr(item, content)
        self.assertEqual(views.parseXlsxRows(BytesIO(rewritten.getvalue())), [['a', 'b']])


class AtomicImportTests(OfficeApiTestBase):
    def test_teacher_import_creates_identity_profile_and_requires_password_change(self):
        ok, errors = views.importTeachers([['TNEW', '新教师', 'new-password']])
        self.assertEqual(len(ok), 1)
        self.assertEqual(errors, [])
        user = User.objects.get(username='TNEW')
        self.assertTrue(user.check_password('new-password'))
        self.assertTrue(user.must_change_password)
        self.assertTrue(user.groups.filter(name='教师').exists())
        self.assertTrue(TeacherProfile.objects.filter(identity__user=user).exists())

    def test_fixed_class_accepts_text_identifier_without_changing_school_number_rules(self):
        ok, errors = views.importStudents([['STUDENT-X', '新学生', 'pw']], self.adminClass)
        self.assertEqual(errors, [])
        self.assertEqual(len(ok), 1)
        profile = StudentProfile.objects.get(identity__school_number='STUDENT-X')
        self.assertEqual(profile.major, self.adminClass.major)
        self.assertEqual(profile.enrollment_year, self.adminClass.enrollment_year)
        self.assertEqual(profile.memberships.get(left_at__isnull=True).admin_class, self.adminClass)

    def test_repeated_teacher_import_preserves_password_name_and_roles(self):
        originalHash = self.teacher.password
        originalName = self.teacher.name
        ok, errors = views.importTeachers([[self.teacher.username, '不应覆盖', 'other-password']])
        self.assertEqual(ok, [])
        self.assertEqual(errors[0][0], 2)
        self.teacher.refresh_from_db()
        self.assertEqual(self.teacher.password, originalHash)
        self.assertEqual(self.teacher.name, originalName)
        self.assertFalse(self.teacher.must_change_password)

    def test_blank_row_is_skipped_and_error_keeps_physical_line_number(self):
        rows = views.parseXlsxRows(xlsx([
            ['工号', '姓名', '密码'], [], ['TNEW', '甲', 'pw'], [], ['TNEW2', '', 'pw']]))
        ok, errors = views.importTeachers(rows)
        self.assertEqual(len(ok), 1)
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0][0], 5)

    def test_invalid_fields_are_rejected_before_any_account_creation(self):
        for row in (['T' * 33, '姓名', 'pw'], ['bad name', '姓名', 'pw'],
                    ['TNEW', '名' * 51, 'pw'], ['TNEW', '姓名', 'p' * 1025],
                    ['TNEW', '\ud800', 'pw'], ['TNEW', '姓名', '\ud800'],
                    ['TNEW', '姓名\x00', 'pw'], ['TNEW', '', 'pw']):
            with self.subTest(row=str(row)[:50]):
                before = User.objects.count()
                ok, errors = views.importTeachers([row])
                self.assertEqual(ok, [])
                self.assertEqual(len(errors), 1)
                self.assertEqual(User.objects.count(), before)

    def test_normalized_username_is_checked_for_duplicate_before_creating(self):
        User.objects.create_user(username='T123', password='original')
        ok, errors = views.importTeachers([['Ｔ１２３', '重复', 'replacement']])
        self.assertEqual(ok, [])
        self.assertEqual(len(errors), 1)
        self.assertTrue(User.objects.get(username='T123').check_password('original'))

    def test_teacher_profile_failure_rolls_back_account_identity_and_group(self):
        actual = TeacherProfile.objects.create

        def create(**kwargs):
            if kwargs['identity'].school_number == 'TFAIL':
                raise IntegrityError('SECRET_SQL_DETAIL')
            return actual(**kwargs)

        with mock.patch.object(TeacherProfile.objects, 'create', side_effect=create):
            ok, errors = views.importTeachers([
                ['TFAIL', '失败', 'SECRET_PASSWORD'], ['TOK', '成功', 'pw']])
        self.assertEqual(len(ok), 1)
        self.assertEqual(len(errors), 1)
        self.assertNotIn('SECRET_SQL_DETAIL', str(errors))
        self.assertNotIn('SECRET_PASSWORD', str(errors))
        self.assertFalse(User.objects.filter(username='TFAIL').exists())
        self.assertFalse(SchoolIdentity.objects.filter(school_number='TFAIL').exists())
        self.assertTrue(User.objects.filter(username='TOK').exists())

    def test_fixed_class_membership_failure_rolls_back_whole_student_row(self):
        with mock.patch.object(views, 'enrollStudent', side_effect=IntegrityError('private detail')):
            ok, errors = views.importStudents([['SFAIL', '失败学生', 'pw']], self.adminClass)
        self.assertEqual(ok, [])
        self.assertEqual(len(errors), 1)
        self.assertFalse(User.objects.filter(username='SFAIL').exists())
        self.assertFalse(StudentProfile.objects.filter(identity__school_number='SFAIL').exists())

    def test_auto_student_failure_also_rolls_back_new_organization_records(self):
        before = (Department.objects.count(), Major.objects.count(), AdminClass.objects.count())
        with mock.patch.object(views, 'enrollStudent', side_effect=IntegrityError('private detail')):
            ok, errors = views.importStudentsAuto([['20231001', '失败学生', 'pw']])
        self.assertEqual(ok, [])
        self.assertEqual(len(errors), 1)
        self.assertEqual((Department.objects.count(), Major.objects.count(),
                          AdminClass.objects.count()), before)
        self.assertFalse(User.objects.filter(username='20231001').exists())
        self.assertFalse(ClassMembership.objects.filter(
            student_profile__identity__school_number='20231001').exists())

    def test_identity_conflict_with_different_login_does_not_leave_partial_account(self):
        existing = User.objects.create_user(username='different-login')
        SchoolIdentity.objects.create(user=existing, school_number='TCONFLICT',
                                       identity_type=SchoolIdentity.IDENTITY_TEACHER)
        ok, errors = views.importTeachers([['TCONFLICT', '新教师', 'pw']])
        self.assertEqual(ok, [])
        self.assertEqual(len(errors), 1)
        self.assertFalse(User.objects.filter(username='TCONFLICT').exists())
        self.assertEqual(SchoolIdentity.objects.filter(school_number='TCONFLICT').count(), 1)

    def test_course_rows_validate_lengths_and_continue_after_invalid_row(self):
        ok, errors = views.importCourses([
            ['X' * 21, '课程'], ['NEW', '名' * 51], ['INVALID', '\ud800'], ['OKNEW', '课程']])
        self.assertEqual(ok, ['课程'])
        self.assertEqual([line for line, _ in errors], [2, 3, 4])
        self.assertEqual(Course.objects.filter(code='OKNEW').count(), 1)

    def test_course_teacher_import_requires_assignment_before_creating_any_user(self):
        self._grant(self.bareUser, 'class_manage')
        with self.assertRaises(PermissionDenied):
            views.importTeachersToCourse([['TNEW', '教师', 'pw']], self.cc, actor=self.bareUser)
        self.assertFalse(User.objects.filter(username='TNEW').exists())

    def test_course_teacher_import_assigns_pool_and_existing_classes_atomically(self):
        ok, errors = views.importTeachersToCourse(
            [['TNEW', '教师', 'pw']], self.cc, actor=self.adminUser)
        self.assertEqual(errors, [])
        self.assertEqual(len(ok), 1)
        teacher = User.objects.get(username='TNEW')
        self.assertTrue(self.cc.teachers.filter(pk=teacher.pk).exists())
        self.assertTrue(self.teachingClass.teachers.filter(pk=teacher.pk).exists())

    def test_course_assignment_failure_rolls_back_created_teacher_and_pool(self):
        through = self.teachingClass.teachers.through
        from django.db.models.signals import m2m_changed

        def reject(sender, action, **kwargs):
            if action == 'pre_add':
                raise IntegrityError('SECRET_RELATION_DETAIL')

        m2m_changed.connect(reject, sender=through, weak=False)
        try:
            ok, errors = views.importTeachersToCourse(
                [['TFAIL', '失败教师', 'pw']], self.cc, actor=self.adminUser)
        finally:
            m2m_changed.disconnect(reject, sender=through)
        self.assertEqual(ok, [])
        self.assertEqual(len(errors), 1)
        self.assertFalse(User.objects.filter(username='TFAIL').exists())
        self.assertFalse(self.cc.teachers.filter(username='TFAIL').exists())
        self.assertFalse(SchoolIdentity.objects.filter(school_number='TFAIL').exists())

    def test_invalid_auto_number_does_not_create_organizations(self):
        before = (Department.objects.count(), Major.objects.count(), AdminClass.objects.count())
        ok, errors = views.importStudentsAuto([['not-a-number', '姓名', 'pw']])
        self.assertEqual(ok, [])
        self.assertEqual(len(errors), 1)
        self.assertEqual((Department.objects.count(), Major.objects.count(),
                          AdminClass.objects.count()), before)


class ImportBoundaryApiTests(OfficeApiTestBase):
    def test_bad_file_is_400_for_all_import_endpoints(self):
        for view, data, kwargs in (
                (ImportView, {'type': 'teacher'}, {}),
                (CourseCollegeView, {'action': 'importStudents'},
                 {'courseId': self.course.pk, 'departmentId': self.department.pk}),
                (AdminClassDetailView, {'action': 'importStudents'},
                 {'adminClassId': self.adminClass.pk})):
            with self.subTest(view=view.__name__):
                data = {**data, 'file': SimpleUploadedFile('bad.xlsx', b'BAD_FILE_PRIVATE')}
                response = self._call(view, self.adminUser, method='post', data=data,
                                      format='multipart', **kwargs)
                self.assertEqual(response.status_code, 400, response.data)
                self.assertNotIn('BAD_FILE_PRIVATE', str(response.data))

    def test_api_errors_use_physical_rows_and_successful_rows_still_commit(self):
        response = self._call(ImportView, self.adminUser, method='post', data={
            'type': 'teacher', 'file': xlsx([
                ['工号', '姓名', '密码'], [], ['TGOOD', '成功', 'pw'], [], ['TBAD', '', 'pw']])},
            format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['successCount'], 1)
        self.assertEqual(response.data['errors'][0]['line'], 5)
        self.assertTrue(User.objects.filter(username='TGOOD').exists())
        self.assertFalse(User.objects.filter(username='TBAD').exists())

    def test_auto_student_failure_rolls_back_new_records_and_next_row_succeeds(self):
        actual = views.enrollStudent

        def enroll(profile, adminClass):
            if profile.identity.school_number == '20231001':
                raise IntegrityError('SECRET_DATABASE_DETAIL')
            return actual(profile, adminClass)

        with mock.patch.object(views, 'enrollStudent', side_effect=enroll):
            response = self._call(ImportView, self.adminUser, method='post', data={
                'type': 'student', 'file': xlsx([
                    ['学号', '姓名', '密码'], ['20231001', '失败', 'SECRET_PASSWORD'],
                    ['20231101', '成功', 'pw']])}, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['successCount'], 1)
        self.assertEqual(response.data['errors'][0]['line'], 2)
        self.assertNotIn('SECRET_DATABASE_DETAIL', str(response.data))
        self.assertNotIn('SECRET_PASSWORD', str(response.data))
        self.assertFalse(User.objects.filter(username='20231001').exists())
        self.assertFalse(Department.objects.filter(name='计算机学院').exists())
        self.assertTrue(User.objects.filter(username='20231101').exists())
        self.assertTrue(Department.objects.filter(name='外国语学院').exists())

    def test_existing_account_password_is_not_replaced_through_api(self):
        original = self.teacher.password
        response = self._call(ImportView, self.adminUser, method='post', data={
            'type': 'teacher', 'file': xlsx([
                ['工号', '姓名', '密码'], [self.teacher.username, '覆盖姓名', 'replacement']])},
            format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['successCount'], 0)
        self.teacher.refresh_from_db()
        self.assertEqual(self.teacher.password, original)

    def test_field_limits_return_row_errors_without_creating_users(self):
        response = self._call(ImportView, self.adminUser, method='post', data={
            'type': 'teacher', 'file': xlsx([
                ['工号', '姓名', '密码'], ['T' * 33, '姓名', 'pw'],
                ['TLONGNAME', '名' * 51, 'pw'], ['TLONGPW', '姓名', 'p' * 1025]])},
            format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['successCount'], 0)
        self.assertEqual([error['line'] for error in response.data['errors']], [2, 3, 4])
        self.assertFalse(User.objects.filter(username__in=['T' * 33, 'TLONGNAME', 'TLONGPW']).exists())

    def test_course_teacher_import_rejects_without_assignment_permission_before_parsing(self):
        self._grant(self.bareUser, 'class_manage')
        self._grant(self.bareUser, 'course_manage')
        with mock.patch.object(views, 'load_workbook') as loader:
            response = self._call(CourseCollegeView, self.bareUser, method='post', data={
                'action': 'importTeachers', 'file': SimpleUploadedFile('bad.xlsx', b'bad')},
                format='multipart', courseId=self.course.pk, departmentId=self.department.pk)
        self.assertEqual(response.status_code, 403, response.data)
        loader.assert_not_called()
        self.assertFalse(User.objects.filter(username='TNEW').exists())

    def test_course_teacher_import_with_assignment_permission_connects_all_relations(self):
        response = self._call(CourseCollegeView, self.adminUser, method='post', data={
            'action': 'importTeachers',
            'file': xlsx([['工号', '姓名', '密码'], ['TNEW', '新教师', 'pw']])},
            format='multipart', courseId=self.course.pk, departmentId=self.department.pk)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['successCount'], 1)
        self.assertTrue(self.cc.teachers.filter(username='TNEW').exists())
        self.assertTrue(self.teachingClass.teachers.filter(username='TNEW').exists())

    def test_fixed_class_api_keeps_text_identifiers_and_builds_current_membership(self):
        response = self._call(AdminClassDetailView, self.adminUser, method='post', data={
            'action': 'importStudents',
            'file': xlsx([['学号', '姓名', '密码'], ['STUDENT-X', '学生', 'pw']])},
            format='multipart', adminClassId=self.adminClass.pk)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['successCount'], 1)
        self.assertTrue(ClassMembership.objects.filter(
            student_profile__identity__school_number='STUDENT-X',
            admin_class=self.adminClass, left_at__isnull=True).exists())

    def test_html_import_bad_file_redirects_with_safe_message(self):
        from django.urls import reverse
        self.client.force_login(self.adminUser)
        response = self.client.post(reverse('office_import'), {
            'type': 'teacher', 'file': SimpleUploadedFile('bad.xlsx', b'PRIVATE_FILE_CONTENT')})
        self.assertEqual(response.status_code, 302)
        messages = [str(message) for message in get_messages(response.wsgi_request)]
        self.assertTrue(any('xlsx' in message for message in messages))
        self.assertNotIn('PRIVATE_FILE_CONTENT', str(messages))
