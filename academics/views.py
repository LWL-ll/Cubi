import io
from datetime import date
from zipfile import BadZipFile, ZipFile

from django.contrib import messages
from django.contrib.auth.models import Group
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import DataError, IntegrityError, transaction
from django.http import Http404
from django.shortcuts import get_object_or_404, redirect, render

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from users.models import SchoolIdentity, StudentProfile, TeacherProfile, User
from users.permissions import hasPerm, requirePerm

from .models import (AdminClass, ClassMembership, Course, CourseCollege,
                     Department, Major, TeachingClass)
from .services import (createTeachingClass, deleteCourseCollege, enrollStudent,
                       removeCandidateAdminClass)
from .student_no import parseStudentNo

# 当前学期：教学班组班默认值（后续做学期管理后再接学期表）
CURRENT_TERM = '2026-2027-1'
MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_IMPORT_ROWS = 10000
MAX_IMPORT_PASSWORD_LENGTH = 1024


# ---------- 工具函数 ----------

def cellToStr(cell):
    """Excel 单元格转字符串：数字型学号/工号防止变成 20231001.0 这种浮点格式"""
    if cell is None:
        return ''
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def parseXlsxRows(fileObj):
    """解析上传的 xlsx 文件，返回数据行列表

    约定：所有导入 xlsx 的第一行必须是表头，解析时直接跳过；
    列约定：A 列编号（学号/工号/课程编号）、B 列姓名/课程名称、C 列初始密码；
    返回的每行对应物理行号从 2 开始（enumerate start=2 的报错行号正好对应）
    """
    if getattr(fileObj, 'size', 0) > MAX_XLSX_BYTES:
        raise ValidationError('xlsx 文件不能超过 10MB')
    workbook = None
    try:
        data = fileObj.read(MAX_XLSX_BYTES + 1)
        if not data or len(data) > MAX_XLSX_BYTES:
            raise ValidationError('xlsx 文件为空或超过 10MB')
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValidationError('xlsx 解压后的内容不能超过 64MB')
            if any(item.flag_bits & 1 for item in archive.infolist()):
                raise ValidationError('不支持加密的 xlsx 文件')
        workbook = load_workbook(io.BytesIO(data), read_only=True)
        sheet = workbook.active
        if sheet is None:
            raise ValidationError('xlsx 中没有可读取的工作表')
        # 不信任表格声明的行列范围，按实际 XML 行读取并及时截断。
        sheet.reset_dimensions()
        rows = []
        for lineNo, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if lineNo > MAX_IMPORT_ROWS + 1:
                raise ValidationError('xlsx 最多允许 10000 行数据（含空白行）')
            if any(cellToStr(cell) for cell in row[3:]):
                raise ValidationError(f'第 {lineNo} 行只能在 A～C 列填写内容')
            if lineNo > 1:
                # 空行保留位置，由导入循环跳过，逐行报错仍对应 Excel 物理行号。
                rows.append([cellToStr(cell) for cell in row[:3]])
        return rows
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError,
            TypeError, SyntaxError, EOFError, RuntimeError, NotImplementedError):
        raise ValidationError('无法解析 xlsx 文件，请检查文件格式与内容') from None
    finally:
        if workbook is not None:
            workbook.close()


def getGroup(name):
    """按中文组名取分组对象"""
    return Group.objects.get(name=name)


def isTeacherUser(user):
    """判断用户是否为教师分组的账号"""
    return user.groups.filter(name='教师').exists()


def findTeacherByNo(teacherNo):
    """按工号查教师账号；不存在/非教师分组/无教师档案 返回错误原因

    严格化（2026-10-18）：教师必须有教师身份档案
    （SchoolIdentity + TeacherProfile）才能进入任课流程，
    手工在 admin 后台建的裸分组账号不允许直接添加。
    """
    user = User.objects.filter(username=teacherNo).first()
    if user is None:
        return None, f'工号 {teacherNo} 的教师不存在，请先在批量导入中导入教师'
    if not isTeacherUser(user):
        return None, f'工号 {teacherNo} 不是教师账号'
    if not user.identities.filter(
            identity_type=SchoolIdentity.IDENTITY_TEACHER,
            teacher_profile__isnull=False).exists():
        return None, f'工号 {teacherNo} 缺少教师档案，请通过批量导入创建教师'
    return user, ''


def ensureDepartment(name):
    """按名称取学院，不存在则创建（编码 D + 递增两位数字）

    学院由学号解析自动生成，无单独导入入口；
    取现有最大数字编码 + 1，避免删除学院后编码重复。
    """
    department = Department.objects.filter(name=name).first()
    if department is not None:
        return department
    codes = Department.objects.filter(code__startswith='D').values_list(
        'code', flat=True)
    nextNo = max((int(code[1:]) for code in codes
                  if code[1:].isdigit()), default=0) + 1
    return Department.objects.create(name=name, code=f'D{nextNo:02d}')


def ensureMajor(department, majorCode, majorName):
    """按（学院, 专业码）取专业，不存在则创建"""
    major, _ = Major.objects.get_or_create(
        department=department, code=majorCode, defaults={'name': majorName})
    return major


def ensureAdminClass(major, enrollmentYear, clazzName, classCode=''):
    """按（专业, 年级, 名称）取行政班，不存在则创建"""
    adminClass, _ = AdminClass.objects.get_or_create(
        major=major, enrollment_year=enrollmentYear, name=clazzName,
        defaults={'code': classCode})
    return adminClass


def createStudentAccount(username, name, password, adminClass, studentGroup):
    """建学生账号 + 学校身份 + 学生资料 + 班级成员关系（V2 三连）"""
    user = User.objects.create_user(username=username, password=password)
    user.name = name
    user.groups.add(studentGroup)
    # 导入账号初始密码在 xlsx 中明文流转，标记首次登录须改密
    user.must_change_password = True
    user.save()
    identity = SchoolIdentity.objects.create(
        user=user, school_number=username,
        identity_type=SchoolIdentity.IDENTITY_STUDENT)
    # 入学年份优先取班级年级，班级没年级时退化学号前 4 位
    enrollmentYear = adminClass.enrollment_year
    if enrollmentYear is None:
        enrollmentYear = (int(username[:4])
                          if username[:4].isdigit() else date.today().year)
    profile = StudentProfile.objects.create(
        identity=identity, major=adminClass.major,
        enrollment_year=enrollmentYear)
    # 统一走服务层入班（一人一班约束收敛在 services.enrollStudent）
    enrollStudent(profile, adminClass)
    return user


def createTeacherAccount(username, name, password, teacherGroup):
    """建教师账号 + 学校身份 + 教师资料（V2）"""
    user = User.objects.create_user(username=username, password=password)
    user.name = name
    user.groups.add(teacherGroup)
    # 导入账号初始密码在 xlsx 中明文流转，标记首次登录须改密
    user.must_change_password = True
    user.save()
    identity = SchoolIdentity.objects.create(
        user=user, school_number=username,
        identity_type=SchoolIdentity.IDENTITY_TEACHER)
    TeacherProfile.objects.create(identity=identity)
    return user


def validateImportText(value, label, maxLength):
    if not isinstance(value, str) or not value.strip() or len(value) > maxLength:
        raise ValidationError(f'{label}须为 1～{maxLength} 字的文本')
    try:
        value.encode('utf-8')
    except UnicodeEncodeError:
        raise ValidationError(f'{label}含无效字符') from None
    if '\x00' in value:
        raise ValidationError(f'{label}含无效字符')
    return value.strip()


def validateAccountRow(cells, label):
    if len(cells) < 3:
        raise ValidationError(f'{label}、姓名、密码均不能为空')
    username = validateImportText(cells[0], label, 32)
    username = User.normalize_username(username)
    username = validateImportText(username, label, 32)
    User._meta.get_field('username').run_validators(username)
    name = validateImportText(cells[1], '姓名', 50)
    password = validateImportText(cells[2], '初始密码', MAX_IMPORT_PASSWORD_LENGTH)
    return username, name, password


def runImportRows(rows, importRow):
    """一行一事务；失败行完整回滚，成功行保留，错误不暴露数据库详情。"""
    okList, errList = [], []
    for lineNo, cells in enumerate(rows, start=2):
        if not any(cells):
            continue
        try:
            with transaction.atomic():
                imported = importRow(cells)
            okList.append(imported)
        except ValidationError as err:
            errList.append((lineNo, '；'.join(err.messages)))
        except IntegrityError:
            errList.append((lineNo, '编号或关联数据冲突，本行未导入'))
        except DataError:
            errList.append((lineNo, '导入数据格式无效，本行未导入'))
    return okList, errList


def importCourses(rows):
    """按行导入课程：每行 [课程编号, 课程名称]，返回 (成功列表, 错误列表)。"""
    def importRow(cells):
        if len(cells) < 2:
            raise ValidationError('课程编号和名称不能为空')
        code = validateImportText(cells[0], '课程编号', 20)
        name = validateImportText(cells[1], '课程名称', 50)
        if Course.objects.filter(code=code).exists():
            raise ValidationError(f'课程编号 {code} 已存在')
        Course.objects.create(code=code, name=name)
        return name
    return runImportRows(rows, importRow)


def importTeachers(rows):
    """按行导入教师账号：每行 [工号, 姓名, 初始密码]，建身份与资料并加入教师分组"""
    teacherGroup = getGroup('教师')
    def importRow(cells):
        username, name, password = validateAccountRow(cells, '工号')
        if User.objects.filter(username=username).exists():
            raise ValidationError(f'工号 {username} 已存在')
        createTeacherAccount(username, name, password, teacherGroup)
        return f'{name}（{username}）'
    return runImportRows(rows, importRow)


def importTeachersToCourse(rows, cc, *, actor):
    """导入教师并加入课程×学院的教师池，同时自动加入已组建的教学班

    每行 [工号, 姓名, 初始密码]，重复工号按行报错跳过。
    """
    if not hasPerm(actor, 'class:assign_teacher'):
        raise PermissionDenied('导入课程教师并分配任课需要 class:assign_teacher 权限')
    teacherGroup = getGroup('教师')
    def importRow(cells):
        username, name, password = validateAccountRow(cells, '工号')
        if User.objects.filter(username=username).exists():
            raise ValidationError(f'工号 {username} 已存在')
        user = createTeacherAccount(username, name, password, teacherGroup)
        # 加入本课程教师池，并自动分配到本课程已组建的每个教学班
        cc.teachers.add(user)
        for teachingClass in cc.teaching_classes.all():
            teachingClass.teachers.add(user)
        return f'{name}（{username}）'
    return runImportRows(rows, importRow)


def importStudents(rows, adminClass):
    """按行导入学生账号到指定行政班：每行 [学号, 姓名, 初始密码]"""
    studentGroup = getGroup('学生')
    def importRow(cells):
        username, name, password = validateAccountRow(cells, '学号')
        if User.objects.filter(username=username).exists():
            raise ValidationError(f'学号 {username} 已存在')
        createStudentAccount(username, name, password, adminClass, studentGroup)
        return f'{name}（{username}）'
    return runImportRows(rows, importRow)


def importStudentsAuto(rows):
    """按学号格式自动分班导入学生：解析出学院/专业/班级后自动创建并按班入学生"""
    studentGroup = getGroup('学生')
    def importRow(cells):
        username, name, password = validateAccountRow(cells, '学号')
        # 按学号格式解析学院/专业/班级（格式不对直接报错，不落库）
        try:
            departmentName, majorCode, majorName, clazzName = parseStudentNo(username)
        except ValueError as err:
            raise ValidationError(str(err)) from None
        if User.objects.filter(username=username).exists():
            raise ValidationError(f'学号 {username} 已存在')
        # 学院/专业/班级不存在则按学号规则自动创建（班级编码 = 专业码 + 班序）
        department = ensureDepartment(departmentName)
        major = ensureMajor(department, majorCode, majorName)
        adminClass = ensureAdminClass(
            major, int(username[0:4]), clazzName, classCode=username[4:8])
        createStudentAccount(username, name, password, adminClass, studentGroup)
        return f'{name}（{username} → {clazzName}）'
    return runImportRows(rows, importRow)


def showImportResult(request, tip, okList, errList):
    """统一的导入结果提示：成功数 + 逐行失败原因"""
    messages.success(request, f'{tip}导入完成：成功 {len(okList)} 条')
    for lineNo, reason in errList:
        messages.warning(request, f'第 {lineNo} 行跳过：{reason}')


def setCourseCollegeLeader(cc, teacher):
    """指定组长：模型层统一同步新旧组长的分组及教师池。"""
    cc.leader = teacher
    cc.save(update_fields=['leader'])


def unsetCourseCollegeLeader(cc):
    """撤销课程组组长：若他不再担任任何课程组组长，则移出「课程组」分组"""
    cc.leader = None
    cc.save(update_fields=['leader'])


# ---------- 页面视图 ----------

@requirePerm('class:manage')
def renderImportPage(request):
    """批量导入页：导入课程 / 导入教师 / 导入班级学生（xlsx 上传）

    学生导入按学号格式自动分配学院、专业和班级，不再手动选择目标班级。
    """
    if request.method == 'POST':
        importType = request.POST.get('type', '')
        csvFile = request.FILES.get('file')
        if not csvFile:
            messages.error(request, '请先选择要上传的 xlsx 文件')
            return redirect('office_import')
        # 动作级权限校验：导入课程需要 course:manage，导入教师/学生需要 class:manage
        needCode = 'course:manage' if importType == 'course' else 'class:manage'
        if not hasPerm(request.user, needCode):
            messages.error(request, '没有权限执行该导入')
            return redirect('office_import')
        importers = {'course': (importCourses, '课程'),
                     'teacher': (importTeachers, '教师'),
                     'student': (importStudentsAuto, '学生')}
        if importType not in importers:
            messages.error(request, '未知的导入类型')
            return redirect('office_import')
        importer, label = importers[importType]
        try:
            okList, errList = importer(parseXlsxRows(csvFile))
        except ValidationError as err:
            messages.error(request, '；'.join(err.messages))
        else:
            showImportResult(request, label, okList, errList)
        return redirect('office_import')
    return render(request, 'office/import.html')


@requirePerm('course:manage')
def renderCourseList(request):
    """课程列表：展示已导入的全部课程及开设学院数"""
    courses = Course.objects.all().prefetch_related('course_colleges__department')
    return render(request, 'office/course_list.html', {'courses': courses})


@requirePerm('course:manage')
def renderCourseDetail(request, courseId):
    """课程设置页：展示已开设学院列表，下拉添加学院进入 课程×学院 设置"""
    course = get_object_or_404(Course, id=courseId)
    if request.method == 'POST':
        # 添加/移除学院属于课程分配：需要 course:assign_class
        if not hasPerm(request.user, 'course:assign_class'):
            messages.error(request, '没有权限执行该操作')
            return redirect('office_course_detail', courseId=courseId)
        action = request.POST.get('action', '')
        if action == 'addDepartment':
            department = get_object_or_404(Department, id=request.POST.get('departmentId'))
            cc, created = CourseCollege.objects.get_or_create(
                course=course, department=department)
            messages.success(request, f'已为课程开设学院 {department.name}')
        elif action == 'removeDepartment':
            department = get_object_or_404(Department, id=request.POST.get('departmentId'))
            cc = get_object_or_404(CourseCollege, course=course, department=department)
            try:
                deleteCourseCollege(cc)
            except ValidationError as err:
                messages.error(request, '；'.join(err.messages))
            else:
                messages.success(request, f'已移除学院 {department.name}')
        return redirect('office_course_detail', courseId=courseId)

    ccList = CourseCollege.objects.filter(course=course).select_related(
        'department', 'leader').prefetch_related('admin_classes', 'teachers')
    # 可添加的学院 = 还没有为这门课开设的学院
    availableDepartments = Department.objects.exclude(course_colleges__course=course)
    return render(request, 'office/course_detail.html', {
        'course': course, 'ccList': ccList,
        'availableDepartments': availableDepartments})


@requirePerm('course:manage')
def renderCourseCollege(request, courseId, departmentId):
    """课程×学院 设置页：课程组组长、本课程教师池、本课程行政班、教学班"""
    cc = get_object_or_404(CourseCollege, course_id=courseId, department_id=departmentId)
    if request.method == 'POST':
        action = request.POST.get('action', '')
        # 每个动作细分权限码校验
        if action in ('setLeader', 'unsetLeader'):
            needCode = 'user:assign_role'
        elif action in ('addAdminClass', 'removeAdminClass', 'createTeachingClass'):
            needCode = 'course:assign_class'
        else:
            needCode = 'class:manage'  # importTeachers / importStudents
        if not hasPerm(request.user, needCode):
            messages.error(request, '没有权限执行该操作')
            return redirect('office_course_college',
                            courseId=courseId, departmentId=departmentId)

        if action == 'setLeader':
            teacherNo = request.POST.get('teacherNo', '').strip()
            teacher, errMsg = findTeacherByNo(teacherNo)
            if errMsg:
                messages.error(request, errMsg)
            else:
                setCourseCollegeLeader(cc, teacher)
                messages.success(request,
                    f'已设置 {teacher.name or teacher.username}（{teacherNo}）为课程组组长，'
                    '并加入「课程组」分组')
        elif action == 'unsetLeader':
            leader = cc.leader
            unsetCourseCollegeLeader(cc)
            messages.success(request, f'已撤销组长 {leader.name or leader.username}')
        elif action == 'importTeachers':
            csvFile = request.FILES.get('file')
            if not csvFile:
                messages.error(request, '请先选择要上传的 xlsx 文件')
            else:
                try:
                    if not hasPerm(request.user, 'class:assign_teacher'):
                        raise PermissionDenied('没有分配教师到教学班的权限')
                    okList, errList = importTeachersToCourse(
                        parseXlsxRows(csvFile), cc, actor=request.user)
                except ValidationError as err:
                    messages.error(request, '；'.join(err.messages))
                except PermissionDenied as err:
                    messages.error(request, str(err))
                else:
                    showImportResult(request, '本课程教师', okList, errList)
        elif action == 'addAdminClass':
            adminClass = get_object_or_404(AdminClass, id=request.POST.get('adminClassId'))
            # 只允许添加本学院的班级（下拉框里本来就是本学院的，这里再兜底校验一次）
            if adminClass.major.department_id != cc.department_id:
                messages.error(request,
                    f'班级 {adminClass.name} 不属于 {cc.department.name}，无法添加')
            else:
                cc.admin_classes.add(adminClass)
                messages.success(request, f'已添加班级 {adminClass.name}')
        elif action == 'removeAdminClass':
            adminClass = get_object_or_404(AdminClass, id=request.POST.get('adminClassId'))
            try:
                removeCandidateAdminClass(cc, adminClass)
            except ValidationError as err:
                messages.error(request, '；'.join(err.messages))
            else:
                messages.success(request, f'已移除班级 {adminClass.name}（学生和教师账号保留）')
        elif action == 'createTeachingClass':
            name = request.POST.get('name', '').strip()
            if not name:
                messages.error(request, '教学班名称不能为空')
            elif TeachingClass.objects.filter(
                    course_college=cc, term=CURRENT_TERM, name=name).exists():
                messages.error(request, f'同学期已存在同名教学班 {name}')
            else:
                adminClassIds = request.POST.getlist('adminClassIds')
                adminClasses = cc.admin_classes.filter(id__in=adminClassIds)
                if not adminClasses:
                    messages.error(request, '请至少选择一个组成班级')
                else:
                    try:
                        createTeachingClass(cc, name, CURRENT_TERM, adminClasses)
                    except ValidationError as err:
                        messages.error(request, '；'.join(err.messages))
                    else:
                        messages.success(request,
                            f'已创建教学班 {name}（{adminClasses.count()} 个班级，已继承教师池）')
        elif action == 'importStudents':
            csvFile = request.FILES.get('file')
            if not csvFile:
                messages.error(request, '请先选择要上传的 xlsx 文件')
            else:
                try:
                    okList, errList = importStudentsAuto(parseXlsxRows(csvFile))
                except ValidationError as err:
                    messages.error(request, '；'.join(err.messages))
                else:
                    showImportResult(request, '学生', okList, errList)
        return redirect('office_course_college', courseId=courseId, departmentId=departmentId)

    adminClasses = cc.admin_classes.prefetch_related('memberships')
    teachers = cc.teachers.all()
    teachingClasses = cc.teaching_classes.prefetch_related(
        'teachers', 'admin_classes')
    # 可添加的班级 = 本学院下还没关联到本课程的行政班
    availableAdminClasses = AdminClass.objects.filter(
        major__department=cc.department).exclude(course_colleges=cc)
    return render(request, 'office/course_college.html', {
        'cc': cc, 'course': cc.course, 'department': cc.department,
        'adminClasses': adminClasses, 'teachers': teachers,
        'teachingClasses': teachingClasses,
        'availableAdminClasses': availableAdminClasses})


@requirePerm('class:manage')
def renderClassDetail(request, courseId, departmentId, adminClassId):
    """行政班详情（页面）：查看学生与所属教学班

    必须从 课程×学院 入口进入（路径自带课程上下文）；
    教师分配在 V2 下挂教学班，移步教学班详情页操作。
    """
    cc = get_object_or_404(CourseCollege, course_id=courseId, department_id=departmentId)
    adminClass = get_object_or_404(AdminClass, id=adminClassId)
    # 校验班级确实属于该课程×学院，防止跨课程访问（404 避免信息泄露）
    if not cc.admin_classes.filter(id=adminClass.id).exists():
        raise Http404('班级不属于该课程学院')

    def back():
        """统一跳回本班级详情页（带课程上下文）"""
        return redirect('office_class_detail', courseId=courseId,
                        departmentId=departmentId, adminClassId=adminClassId)

    if request.method == 'POST':
        action = request.POST.get('action', '')
        if action == 'importStudents':
            if not hasPerm(request.user, 'class:manage'):
                messages.error(request, '没有权限执行该操作')
                return back()
            csvFile = request.FILES.get('file')
            if not csvFile:
                messages.error(request, '请先选择要上传的 xlsx 文件')
            else:
                try:
                    okList, errList = importStudents(parseXlsxRows(csvFile), adminClass)
                except ValidationError as err:
                    messages.error(request, '；'.join(err.messages))
                else:
                    showImportResult(request, '学生', okList, errList)
        else:
            messages.error(request, '未知操作')
        return back()

    # 学生经成员关系推导（只列当前在班成员）
    memberships = adminClass.memberships.select_related(
        'student_profile__identity__user').filter(left_at__isnull=True)
    students = [m.student_profile.identity.user for m in memberships]
    # 本班参与的教学班（当前课程×学院下）
    teachingClasses = adminClass.teaching_classes.filter(course_college=cc)
    return render(request, 'office/class_detail.html', {
        'cc': cc, 'course': cc.course, 'department': cc.department,
        'adminClass': adminClass,
        'students': students, 'teachingClasses': teachingClasses})


@requirePerm('class:manage')
def renderTeachingClassDetail(request, courseId, departmentId, teachingClassId):
    """教学班详情（页面）：教师分配（下拉或工号添加）、组成班级展示

    必须从 课程×学院 入口进入；教师分配需要 class:assign_teacher，
    分配候选来自课程教师池。
    """
    cc = get_object_or_404(CourseCollege, course_id=courseId, department_id=departmentId)
    teachingClass = get_object_or_404(TeachingClass, id=teachingClassId)
    # 教学班必须属于该课程×学院（404 防跨课程探测）
    if teachingClass.course_college_id != cc.id:
        raise Http404('教学班不属于该课程学院')

    def back():
        """统一跳回本教学班详情页（带课程上下文）"""
        return redirect('office_teaching_class_detail', courseId=courseId,
                        departmentId=departmentId, teachingClassId=teachingClassId)

    if request.method == 'POST':
        action = request.POST.get('action', '')
        if action in ('addTeacher', 'removeTeacher'):
            if not hasPerm(request.user, 'class:assign_teacher'):
                messages.error(request, '没有权限执行该操作')
                return back()
            if action == 'addTeacher':
                # 两种方式：下拉框选教师（teacherId）或直接输工号（teacherNo）
                teacherNo = request.POST.get('teacherNo', '').strip()
                if teacherNo:
                    teacher, errMsg = findTeacherByNo(teacherNo)
                    if errMsg:
                        messages.error(request, errMsg)
                    else:
                        # 工号添加的教师同时加入本课程教师池
                        teachingClass.teachers.add(teacher)
                        cc.teachers.add(teacher)
                        messages.success(request,
                            f'已添加教师 {teacher.name or teacher.username}（{teacherNo}）并加入本课程教师池')
                else:
                    teacher = get_object_or_404(User, id=request.POST.get('teacherId'))
                    # 下拉框候选来自教师池，这里再兜底校验一次防止伪造请求
                    if not cc.teachers.filter(id=teacher.id).exists():
                        messages.error(request, '该教师不在本课程教师池中，无法分配')
                    else:
                        teachingClass.teachers.add(teacher)
                        messages.success(request,
                            f'已添加教师 {teacher.name or teacher.username}')
            else:
                teacher = get_object_or_404(User, id=request.POST.get('teacherId'))
                teachingClass.teachers.remove(teacher)
                messages.success(request,
                    f'已移除教师 {teacher.name or teacher.username}')
        else:
            messages.error(request, '未知操作')
        return back()

    # 可分配的教师 = 本课程×学院教师池里还没分配到本教学班的
    availableTeachers = cc.teachers.exclude(
        id__in=teachingClass.teachers.values('id'))
    return render(request, 'office/teaching_class_detail.html', {
        'cc': cc, 'course': cc.course, 'department': cc.department,
        'teachingClass': teachingClass,
        'teachers': teachingClass.teachers.all(),
        'adminClasses': teachingClass.admin_classes.all(),
        'availableTeachers': availableTeachers})
